#!/usr/bin/env python3
"""
Compare the output folder of the Python tool with the output folder of the
PowerShell port, both produced against the SAME Grid (e.g. a lab Grid).

    python compare_runs.py <python_run_dir> <powershell_run_dir>

Exit code 0 = every artifact matches, 1 = at least one difference.

What is compared (and what is deliberately ignored):
  core .csv            every cell; only 'Collected At' ignored; floats compared numerically
  .summary.json        all keys except log_file, collected_at, hashes, format
  capacity             Python's two-sheet .xlsx  vs  PowerShell's two .csv files
  topology .json/.html parsed JSON (key order ignored, list order NOT ignored); the HTML
                       template text outside the embedded JSON must be identical
  license .csv/.json   cell-for-cell / parsed JSON
Needs openpyxl only to read the Python tool's capacity workbook.
"""
import csv
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

results = []   # (name, ok, detail)


def record(name, ok, detail=""):
    results.append((name, ok, detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"\n         {detail}" if detail and not ok else ""))


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def norm(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def cell_equal(a, b):
    if a == b:
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except ValueError:
        return False


def diff_tables(a, b, ignore_cols=(), limit=5):
    """Return a list of human-readable differences between two string tables."""
    out = []
    if not a or not b:
        return ["one table is empty"]
    if a[0] != b[0]:
        out.append(f"headers differ: only-in-python={[h for h in a[0] if h not in b[0]]} "
                   f"only-in-ps={[h for h in b[0] if h not in a[0]]} (or order differs)")
    if len(a) != len(b):
        out.append(f"row count python={len(a)-1} powershell={len(b)-1}")
    for ri, (ra, rb) in enumerate(zip(a[1:], b[1:]), 2):
        if len(ra) != len(rb):
            out.append(f"row {ri}: column count {len(ra)} vs {len(rb)}")
            continue
        for ci, (x, y) in enumerate(zip(ra, rb)):
            if a[0][ci] in ignore_cols:
                continue
            if not cell_equal(x, y):
                out.append(f"row {ri}, col {ci+1} '{a[0][ci]}': python={x!r} powershell={y!r}")
    return out[:limit] + ([f"... and {len(out)-limit} more"] if len(out) > limit else [])


def find(d, suffix):
    hits = sorted(d.glob(f"*{suffix}"))
    # '<base>.csv' must not match '<base>_license_report.csv' etc.
    if suffix == ".csv":
        hits = [h for h in hits if not re.search(r"_(license_report|grid_capacity\w*)\.csv$", h.name)]
    return hits[0] if hits else None


def xlsx_sheet(path, name):
    wb = load_workbook(path, read_only=True, data_only=True)
    return [[norm(c) for c in row] for row in wb[name].iter_rows(values_only=True)]


def main(py_dir, ps_dir):
    py, ps = Path(py_dir), Path(ps_dir)

    # ---- core csv
    a, b = find(py, ".csv"), find(ps, ".csv")
    if not (a and b):
        record("core report .csv", False, f"missing file (python={a}, powershell={b})")
    else:
        d = diff_tables(read_csv(a), read_csv(b), ignore_cols=("Collected At",))
        record("core report .csv (43 columns)", not d, "\n         ".join(d))

    # ---- summary.json
    a, b = find(py, ".summary.json"), find(ps, ".summary.json")
    if not (a and b):
        record("summary.json", False, "missing file")
    else:
        drop = {"log_file", "collected_at", "hashes", "format"}
        ja = {k: v for k, v in json.loads(a.read_text(encoding="utf-8-sig")).items() if k not in drop}
        jb = {k: v for k, v in json.loads(b.read_text(encoding="utf-8-sig")).items() if k not in drop}
        d = [f"{k}: python={ja.get(k)!r} powershell={jb.get(k)!r}" for k in sorted(set(ja) | set(jb)) if ja.get(k) != jb.get(k)]
        record("summary.json", not d, "\n         ".join(d))

    # ---- capacity (python xlsx vs powershell csv pair), only if python produced it
    cap_x = find(py, "_grid_capacity.xlsx")
    if cap_x:
        if load_workbook is None:
            record("capacity report", False, "openpyxl not installed; cannot read the Python workbook")
        else:
            for sheet, suffix in (("Capacity Summary", "_grid_capacity_summary.csv"),
                                  ("Object Counts", "_grid_capacity_objects.csv")):
                b = find(ps, suffix)
                if not b:
                    record(f"capacity: {sheet}", False, f"powershell file *{suffix} missing")
                    continue
                d = diff_tables(xlsx_sheet(cap_x, sheet), read_csv(b))
                record(f"capacity: {sheet}", not d, "\n         ".join(d))

    # ---- topology
    a, b = find(py, "_topology.json"), find(ps, "_topology.json")
    if a or b:
        if not (a and b):
            record("topology.json", False, f"missing file (python={a}, powershell={b})")
        else:
            ja, jb = json.loads(a.read_text(encoding="utf-8-sig")), json.loads(b.read_text(encoding="utf-8-sig"))
            d = []
            for key in ("meta", "zone_counts"):
                if ja[key] != jb[key]:
                    d.append(f"{key}: python={ja[key]} powershell={jb[key]}")
            na, nb = ja["topology"]["nodes"], jb["topology"]["nodes"]
            la, lb = ja["topology"]["links"], jb["topology"]["links"]
            if na != nb:
                d.append(f"nodes differ: python={len(na)} powershell={len(nb)}; "
                         f"only-python={[n['id'] for n in na if n not in nb][:3]} only-ps={[n['id'] for n in nb if n not in na][:3]}")
            if la != lb:
                d.append(f"links differ: python={len(la)} powershell={len(lb)}; "
                         f"first mismatch at index {next((i for i,(x,y) in enumerate(zip(la,lb)) if x!=y), min(len(la),len(lb)))}")
            record("topology.json (nodes, links, order)", not d, "\n         ".join(d))

    a, b = find(py, "_topology.html"), find(ps, "_topology.html")
    if a or b:
        if not (a and b):
            record("topology.html", False, "missing file")
        else:
            pat = re.compile(r"var DDI_DATA = (.*?);\r?\n", re.S)
            ta, tb = a.read_text(encoding="utf-8-sig"), b.read_text(encoding="utf-8-sig")
            ma, mb = pat.search(ta), pat.search(tb)
            d = []
            if not (ma and mb):
                d.append("could not find 'var DDI_DATA = ...;' in one of the files")
            else:
                if json.loads(ma.group(1)) != json.loads(mb.group(1)):
                    d.append("embedded DDI_DATA JSON differs")
                canon = lambda t, m: t.replace(m.group(1), "__JSON__").replace("\r\n", "\n").strip()
                if canon(ta, ma) != canon(tb, mb):
                    d.append("HTML template text outside the JSON differs")
            if "</script>" in (mb.group(1) if mb else ""):
                d.append("embedded JSON contains a raw </script>")
            record("topology.html (template + embedded data)", not d, "\n         ".join(d))

    # ---- license report
    a, b = find(py, "_license_report.csv"), find(ps, "_license_report.csv")
    if a or b:
        if not (a and b):
            record("license_report.csv", False, "missing file")
        else:
            d = diff_tables(read_csv(a), read_csv(b))
            record("license_report.csv", not d, "\n         ".join(d))
    a, b = find(py, "_license_report.json"), find(ps, "_license_report.json")
    if a or b:
        if not (a and b):
            record("license_report.json", False, "missing file")
        else:
            ja, jb = json.loads(a.read_text(encoding="utf-8-sig")), json.loads(b.read_text(encoding="utf-8-sig"))
            record("license_report.json", ja == jb,
                   "parsed JSON differs (e.g. expiry dates or per-member license lists)")

    bad = [r for r in results if not r[1]]
    print(f"\n{len(results) - len(bad)}/{len(results)} checks passed")
    return 1 if bad else 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1], sys.argv[2]))
